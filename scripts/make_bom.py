"""Build the Screen-2 BOM & calculations workbook."""
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FX = 96.2                 # INR per USD, early Oct 2026 (market; Trading Economics)
IN_P, OUT_P = 1.0, 5.0    # Claude Haiku 4.5 list price, USD per million tokens (platform.claude.com/docs pricing)
TOK_IN, TOK_OUT = 1350, 260   # measured averages across 693 real orders (backtest + forward run)
bt = json.load(open("data/backtest/summary.json"))

wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F2A5A"); B = Font(bold=True)


def sheet(title, header, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = H, HF, Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


cost_order_usd = TOK_IN / 1e6 * IN_P + TOK_OUT / 1e6 * OUT_P
cost_order_inr = cost_order_usd * FX

ws = sheet("1 Prototype BOM", ["#", "Item", "Type", "Qty / basis", "Cost (INR)", "Status", "Note / source"], [
    [1, "Developer laptop (existing)", "Hardware", "1", 0, "Owned", "No hardware needs to be bought"],
    [2, "Open court data: Indian High Court Judgments (AWS Open Data)", "Data", "~2,500 order PDFs + metadata", 0, "Used", "CC-BY-4.0; no account needed"],
    [3, "Python, FastAPI, SQLite, Poppler, Tesseract OCR", "Software", "Open source", 0, "Used", "No licence cost"],
    [4, "Claude API (Haiku 4.5) for development + evaluation", "AI usage", "~1,500 order runs so far", 500, "Spent (approx.)", f"Measured ~{TOK_IN}+{TOK_OUT} tokens/order at $1/$5 per M tokens"],
    [5, "Claude API budget: full evaluation, re-runs, live demo", "AI usage", "~10,000 order runs", round(10000 * cost_order_inr), "Budgeted", "Estimate"],
    [6, "Cloud VM for live API demo (2 vCPU, 4 GB) for 3 months", "Hosting", "3 months", 6000, "Budgeted", "Estimate; static demo on GitHub Pages is free"],
    [7, "Domain name", "Hosting", "1 year", 1000, "Budgeted", "Estimate"],
    [8, "Honorarium: blind review by a practising lawyer / legal-cell officer", "Validation", "30-50 obligations", 10000, "Budgeted", "Estimate"],
    [9, "Video recording / editing", "Pitch", "-", 0, "-", "Free tools"],
    ["", "TOTAL prototype", "", "", "=SUM(E2:E10)", "", "Software-only; no hardware BOM"],
], [4, 52, 12, 26, 12, 14, 52], first=True)
ws["B11"].font = ws["E11"].font = B

sheet("2 Unit cost per order", ["Quantity", "Value", "Unit", "Basis"], [
    ["Input tokens per order (measured mean)", TOK_IN, "tokens", "293 backtest + 400 forward-run orders"],
    ["Output tokens per order (measured mean)", TOK_OUT, "tokens", "same"],
    ["Price: input (Claude Haiku 4.5)", IN_P, "USD / million tokens", "platform.claude.com/docs/en/about-claude/pricing"],
    ["Price: output (Claude Haiku 4.5)", OUT_P, "USD / million tokens", "same"],
    ["AI cost per order", "=B2/1e6*B4+B3/1e6*B5", "USD", "formula"],
    ["INR per USD", FX, "INR", "market rate, early Oct 2026"],
    ["AI cost per order", "=B6*B7", "INR", "formula"],
    ["AI cost per order with Batch API (50% off)", "=B8/2", "INR", "Batch API discount 50%"],
    ["OCR + text extraction", 0, "INR", "Open source, runs on the same server"],
    ["Processing time per order (measured)", 3, "seconds", "upload → obligations, live API"],
], [46, 18, 22, 52])

orders = 93749
sheet("3 State deployment (yr)", ["Item", "Qty", "Unit cost (INR)", "Annual cost (INR)", "Basis"], [
    ["AI processing of all orders with a government party (P&H HC, 2025 volume)", orders, round(cost_order_inr, 2), f"=B2*C2", "93,749 orders in 2025 open data with State/UT/Union in title"],
    ["Hosting: 2 app servers + managed database + backups (NIC/MeghRaj or empanelled cloud)", 1, 350000, "=B3*C3", "Estimate"],
    ["Alerts: SMS/WhatsApp/email (≈66,000 tasks × 5 alerts)", 330000, 0.25, "=B4*C4", "Estimate; per-message rate varies"],
    ["Security audit (CERT-In empanelled), amortised", 1, 150000, "=B5*C5", "Estimate, one-time ≈ 1.5 lakh"],
    ["Support and maintenance (part-time engineer)", 1, 500000, "=B6*C6", "Estimate"],
    ["Training of legal-cell nodal officers", 1, 100000, "=B7*C7", "Estimate"],
    ["TOTAL per state / year", "", "", "=SUM(D2:D7)", ""],
    ["Estimated court-ordered tasks tracked per year", 66000, "", "", "62,009 final govt-party orders × 64% with a direction × 1.67 per order (sample-based estimate)"],
    ["Cost per tracked task", "", "", "=D8/B9", "formula"],
], [62, 12, 16, 18, 58])

sheet("4 Calculations", ["Calculation", "Working", "Result"], [
    ["Daily load (one High Court)", "93,749 govt-party orders ÷ ~240 working days", "≈ 390 orders/day"],
    ["Compute time per day", "390 orders × 3 s (sequential)", "≈ 20 minutes/day on one small server"],
    ["Peak throughput", "8 parallel workers × 1 order / 3 s", "≈ 9,600 orders/hour"],
    ["Annual AI cost, one High Court", f"93,749 × ₹{cost_order_inr:.2f}", f"≈ ₹{orders * cost_order_inr:,.0f} (≈ ₹{orders * cost_order_inr / 2:,.0f} with Batch API)"],
    ["Backtest coverage", f"{bt['captured']} of {bt['cases']} original orders behind real contempt cases", f"{bt['captured_pct']}% captured"],
    ["Backtest warning window", f"median (contempt filed − computed deadline), {bt['deadline_before_contempt']} dated cases", f"{bt['median_lead_days']} days"],
    ["Storage", "Text + extractions ≈ 10 KB/order × 93,749", "≈ 1 GB/year (PDFs optional, stay with the court)"],
], [36, 62, 48])

sheet("5 Sources", ["What", "Source"], [
    ["Court orders, metadata", "https://registry.opendata.aws/indian-high-court-judgments/ (CC-BY-4.0)"],
    ["Claude Haiku 4.5 prices, Batch API discount", "https://platform.claude.com/docs/en/about-claude/pricing"],
    ["USD/INR ≈ 96.2 (early Oct 2026)", "https://tradingeconomics.com/india/currency"],
    ["Measured tokens, timings, backtest", "This repository: data/backtest/summary.json, docs/EVALUATION.md"],
    ["All rows marked 'Estimate'", "Team estimates, to be validated in a pilot"],
], [44, 90])
wb.save("submission/screen2/Anupalan_BOM_and_calculations.xlsx")
print(f"cost/order USD {cost_order_usd:.4f} INR {cost_order_inr:.2f}; annual AI ₹{orders*cost_order_inr:,.0f}")
